#!/usr/bin/env python3
"""Generate every format of the AI Governance Checklist from data/checklist.json.

Outputs:
  CHECKLIST.md                 Markdown checklist (GitHub-rendered)
  sheets/checklist.csv         Google Sheets / any spreadsheet import
  sheets/ai-governance-checklist.xlsx   Excel workbook with status dropdowns and live progress
  docs/index.html              Interactive web checklist (GitHub Pages)

Usage: python3 scripts/build.py
"""
import csv
import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parent.parent
DATA = json.loads((ROOT / "data" / "checklist.json").read_text(encoding="utf-8"))
ITEMS = DATA["items"]
STATUSES = DATA["statuses"]
REPO = "https://github.com/caiotk/NxgdAI-Governance-Checklist"
PAGES = "https://caiotk.github.io/NxgdAI-Governance-Checklist/"
RAW_CSV = "https://raw.githubusercontent.com/caiotk/NxgdAI-Governance-Checklist/main/sheets/checklist.csv"


# ---------------------------------------------------------------- Markdown
def build_md():
    out = [
        f"# {DATA['title']}",
        "",
        f"_{DATA['subtitle']}_",
        "",
        f"Version {DATA['version']} · Updated {DATA['updated']} · Published by {DATA['publisher']}",
        "",
        f"**Other formats:** [Interactive web checklist]({PAGES}) · "
        "[Excel workbook](sheets/ai-governance-checklist.xlsx) · "
        "[CSV for Google Sheets](sheets/checklist.csv)",
        "",
        "## Summary",
        "",
        "| # | Control | Area | Owner |",
        "|---|---|---|---|",
    ]
    for it in ITEMS:
        out.append(f"| {it['id']} | [{it['control']}](#{it['id']}--{slug(it['control'])}) | {it['area']} | {it['owner']} |")
    out += ["", "## The checklist", ""]
    for it in ITEMS:
        out += [
            f"### {it['id']} · {it['control']}",
            "",
            f"- [ ] **{it['control']}** — {it['description']}",
            "",
            f"**Evidence an auditor will ask for:** {it['evidence']}",
            "",
            f"**Quick test:** {it['test']}",
            "",
            f"**Suggested owner:** {it['owner']}",
            "",
            "| ISO/IEC 42001 | NIST AI RMF | EU AI Act | Also maps to |",
            "|---|---|---|---|",
            f"| {it['iso42001']} | {it['nist']} | {it['euaiact']} | {it['more']} |",
            "",
        ]
    out += [
        "## How to use status",
        "",
        "Track each control through: " + " → ".join(f"**{s}**" for s in STATUSES[:4]) +
        f", or mark it **{STATUSES[4]}** with a reason. A control only counts as done when it is "
        "**Evidenced** — implemented, and you can show an auditor the proof.",
        "",
        "## Framework mapping note",
        "",
        DATA["mapping_note"],
        "",
        "---",
        "",
        "Strong governance builds trust. Responsible AI drives sustainable innovation.",
        "",
    ]
    (ROOT / "CHECKLIST.md").write_text("\n".join(out), encoding="utf-8")


def slug(s):
    return "".join(c if c.isalnum() else "-" for c in s.lower()).strip("-").replace("--", "-")


# ---------------------------------------------------------------- CSV
CSV_HEADERS = ["#", "Area", "Control", "Description", "Evidence", "Quick test",
               "Suggested owner", "ISO/IEC 42001", "NIST AI RMF", "EU AI Act", "Also maps to",
               "Status", "Assigned to", "Target date", "Evidence link", "Notes"]
STATUS_COL = CSV_HEADERS.index("Status") + 1          # 1-based column of the Status dropdown
STATUS_LETTER = get_column_letter(STATUS_COL)
LAST_LETTER = get_column_letter(len(CSV_HEADERS))


def build_csv():
    p = ROOT / "sheets" / "checklist.csv"
    p.parent.mkdir(exist_ok=True)
    with p.open("w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(CSV_HEADERS)
        for it in ITEMS:
            w.writerow([it["id"], it["area"], it["control"], it["description"], it["evidence"],
                        it["test"], it["owner"], it["iso42001"], it["nist"], it["euaiact"], it["more"],
                        "Not started", "", "", "", ""])


# ---------------------------------------------------------------- XLSX
NAVY, BLUE, GREY = "1F3864", "2E75B6", "595959"
F = "Arial"


def build_xlsx():
    wb = Workbook()
    thin = Side(style="thin", color="D9D9D9")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    wrap = Alignment(wrap_text=True, vertical="top")

    # --- Checklist sheet
    ws = wb.active
    ws.title = "Checklist"
    ws.sheet_view.showGridLines = False
    ws["A1"] = DATA["title"]
    ws["A1"].font = Font(name=F, size=18, bold=True, color=NAVY)
    ws["A2"] = DATA["subtitle"]
    ws["A2"].font = Font(name=F, size=10, italic=True, color=GREY)
    ws["A3"] = (f"Edit only the yellow columns (Status, Assigned to, Target date, Evidence link, Notes). "
                f"Status values come from the dropdown. Version {DATA['version']} · {DATA['updated']}.")
    ws["A3"].font = Font(name=F, size=9, color=GREY)

    hdr_row = 5
    for c, h in enumerate(CSV_HEADERS, 1):
        cell = ws.cell(row=hdr_row, column=c, value=h)
        cell.font = Font(name=F, bold=True, color="FFFFFF", size=10)
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        cell.border = border

    input_fill = PatternFill("solid", fgColor="FFF9DB")
    first = hdr_row + 1
    for r, it in enumerate(ITEMS, first):
        vals = [it["id"], it["area"], it["control"], it["description"], it["evidence"], it["test"],
                it["owner"], it["iso42001"], it["nist"], it["euaiact"], it["more"], "Not started", "", None, "", ""]
        for c, v in enumerate(vals, 1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.font = Font(name=F, size=10, bold=(c == 3), color=(NAVY if c == 3 else "000000"))
            cell.alignment = wrap
            cell.border = border
            if c >= STATUS_COL:
                cell.fill = input_fill
        ws.cell(row=r, column=STATUS_COL + 2).number_format = "yyyy-mm-dd"
    last = first + len(ITEMS) - 1

    # Example row note for the input columns
    ws.cell(row=last + 2, column=1, value=(
        "Example of a filled row: Status = Evidenced · Assigned to = J. Smith (CRO) · "
        "Target date = 2026-12-15 · Evidence link = https://intranet/ai-register · "
        "Notes = Register reviewed by audit committee Nov 2026."))
    ws.cell(row=last + 2, column=1).font = Font(name=F, size=9, italic=True, color=GREY)
    ws.cell(row=last + 3, column=1, value="Framework mapping: " + DATA["mapping_note"])
    ws.cell(row=last + 3, column=1).font = Font(name=F, size=9, italic=True, color=GREY)

    dv = DataValidation(type="list", formula1="=Lists!$A$2:$A$%d" % (len(STATUSES) + 1), allow_blank=True)
    dv.error = "Choose a status from the list."
    ws.add_data_validation(dv)
    dv.add(f"{STATUS_LETTER}{first}:{STATUS_LETTER}{last}")

    # Status colouring
    colours = {"Not started": "F2F2F2", "In progress": "FFE699", "Implemented": "BDD7EE",
               "Evidenced": "C6EFCE", "Not applicable": "D9D9D9"}
    for s, col in colours.items():
        ws.conditional_formatting.add(
            f"{STATUS_LETTER}{first}:{STATUS_LETTER}{last}",
            FormulaRule(formula=[f'${STATUS_LETTER}{first}="{s}"'], fill=PatternFill("solid", fgColor=col)))

    widths = [5, 14, 26, 44, 44, 40, 22, 20, 20, 26, 32, 15, 18, 13, 28, 30]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = ws.cell(row=first, column=4)
    ws.auto_filter.ref = f"A{hdr_row}:{LAST_LETTER}{last}"
    ws.print_title_rows = f"{hdr_row}:{hdr_row}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True

    # --- Lists sheet
    ls = wb.create_sheet("Lists")
    ls["A1"] = "Status"
    ls["A1"].font = Font(name=F, bold=True)
    for i, s in enumerate(STATUSES, 2):
        ls.cell(row=i, column=1, value=s).font = Font(name=F)
    ls.column_dimensions["A"].width = 18

    # --- Dashboard sheet
    db = wb.create_sheet("Dashboard", 0)
    db.sheet_view.showGridLines = False
    db["A1"] = "AI Governance — progress"
    db["A1"].font = Font(name=F, size=18, bold=True, color=NAVY)
    db["A2"] = "Calculated live from the Status column on the Checklist sheet."
    db["A2"].font = Font(name=F, size=10, italic=True, color=GREY)
    rng = f"Checklist!${STATUS_LETTER}${first}:${STATUS_LETTER}${last}"
    db["A4"], db["B4"] = "Status", "Controls"
    for c in ("A4", "B4"):
        db[c].font = Font(name=F, bold=True, color="FFFFFF")
        db[c].fill = PatternFill("solid", fgColor=NAVY)
    for i, s in enumerate(STATUSES, 5):
        db.cell(row=i, column=1, value=f"=Lists!A{i - 3}").font = Font(name=F, color="008000")
        db.cell(row=i, column=2, value=f"=COUNTIF({rng},A{i})").font = Font(name=F)
    tot = 5 + len(STATUSES)
    db.cell(row=tot, column=1, value="Total controls").font = Font(name=F, bold=True)
    db.cell(row=tot, column=2, value=f"=COUNTA({rng})").font = Font(name=F, bold=True)

    db.cell(row=tot + 2, column=1, value="In scope (excludes Not applicable)").font = Font(name=F)
    db.cell(row=tot + 2, column=2, value=f"=B{tot}-B9").font = Font(name=F)
    db.cell(row=tot + 3, column=1, value="Implemented or evidenced").font = Font(name=F)
    db.cell(row=tot + 3, column=2, value="=B7+B8").font = Font(name=F)
    db.cell(row=tot + 4, column=1, value="Coverage (implemented or evidenced)").font = Font(name=F, bold=True)
    c = db.cell(row=tot + 4, column=2, value=f"=IF(B{tot + 2}=0,0,B{tot + 3}/B{tot + 2})")
    c.font = Font(name=F, bold=True, color=BLUE)
    c.number_format = "0%;-0%;-"
    db.cell(row=tot + 5, column=1, value="Audit-ready (evidenced only)").font = Font(name=F, bold=True)
    c = db.cell(row=tot + 5, column=2, value=f"=IF(B{tot + 2}=0,0,B8/B{tot + 2})")
    c.font = Font(name=F, bold=True, color=BLUE)
    c.number_format = "0%;-0%;-"

    db.cell(row=tot + 7, column=1, value="By area").font = Font(name=F, bold=True, color=NAVY, size=12)
    db.cell(row=tot + 8, column=1, value="Area").font = Font(name=F, bold=True)
    db.cell(row=tot + 8, column=2, value="Controls").font = Font(name=F, bold=True)
    db.cell(row=tot + 8, column=3, value="Evidenced").font = Font(name=F, bold=True)
    areas = list(dict.fromkeys(it["area"] for it in ITEMS))
    arng = f"Checklist!$B${first}:$B${last}"
    for i, a in enumerate(areas, tot + 9):
        db.cell(row=i, column=1, value=a).font = Font(name=F)
        db.cell(row=i, column=2, value=f"=COUNTIF({arng},A{i})").font = Font(name=F)
        db.cell(row=i, column=3, value=f'=COUNTIFS({arng},A{i},{rng},"Evidenced")').font = Font(name=F)
    db.column_dimensions["A"].width = 38
    db.column_dimensions["B"].width = 12
    db.column_dimensions["C"].width = 12

    # --- About sheet
    ab = wb.create_sheet("About")
    rows = [
        ("About", None),
        ("Publisher", DATA["publisher"]),
        ("Version", DATA["version"]),
        ("Updated", DATA["updated"]),
        ("Web version", PAGES),
        ("Source repository", REPO),
        ("How to use", "Set a Status for each control on the Checklist sheet. The Dashboard updates itself. "
                       "A control counts as audit-ready only when it is Evidenced."),
        ("Framework mapping", DATA["mapping_note"]),
        ("Licence", "CC BY 4.0 — share and adapt with attribution to NexGuide.AI."),
    ]
    for r, (k, v) in enumerate(rows, 1):
        ab.cell(row=r, column=1, value=k).font = Font(name=F, bold=True, color=NAVY, size=(14 if r == 1 else 10))
        if v:
            cell = ab.cell(row=r, column=2, value=v)
            cell.font = Font(name=F)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    ab.column_dimensions["A"].width = 20
    ab.column_dimensions["B"].width = 90

    wb.properties.title = DATA["title"]
    wb.properties.creator = DATA["publisher"]
    wb.save(ROOT / "sheets" / "ai-governance-checklist.xlsx")


# ---------------------------------------------------------------- HTML
def build_html():
    tpl = (ROOT / "scripts" / "template.html").read_text(encoding="utf-8")
    html = (tpl.replace("__DATA__", json.dumps(DATA, ensure_ascii=False))
               .replace("__REPO__", REPO)
               .replace("__RAW_CSV__", RAW_CSV)
               .replace("__VERSION__", DATA["version"])
               .replace("__UPDATED__", DATA["updated"]))
    (ROOT / "docs").mkdir(exist_ok=True)
    (ROOT / "docs" / "index.html").write_text(html, encoding="utf-8")


if __name__ == "__main__":
    build_md()
    build_csv()
    build_xlsx()
    build_html()
    print("Built CHECKLIST.md, sheets/checklist.csv, sheets/ai-governance-checklist.xlsx, docs/index.html")
